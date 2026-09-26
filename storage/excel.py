from pathlib import Path
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


# ============================================================
# PATHS
# ============================================================

BASE_DIR = Path(__file__).resolve().parent.parent

EXPORT_DIR = BASE_DIR / "data" / "exports"
EXPORT_DIR.mkdir(parents=True, exist_ok=True)


# ============================================================
# HELPERS
# ============================================================

def _to_text(value) -> str:
    """
    Convert lists/dicts/other values into Excel-safe text.
    """

    if value is None:
        return ""

    if isinstance(value, list):
        return "; ".join(
            str(item)
            for item in value
        )

    if isinstance(value, dict):
        return "; ".join(
            f"{key}: {value}"
            for key, value in value.items()
        )

    return str(value)


def _build_row(email_data: dict) -> list:
    """
    Convert one final email/lead record into an Excel row.
    """

    sales_potential = email_data.get(
        "sales_potential",
        {},
    )

    return [
        email_data.get("business_name"),
        email_data.get("category"),
        email_data.get("city"),
        email_data.get("phone"),
        email_data.get("email"),
        email_data.get("website"),

        email_data.get("qualification"),
        email_data.get("qualification_reason"),
        email_data.get("qualification_confidence"),

        email_data.get("research_depth"),
        email_data.get("business_summary"),
        email_data.get("industry"),
        email_data.get("business_activity_level"),

        email_data.get("hiring_signals"),
        email_data.get("expansion_signals"),
        email_data.get("workplace_signals"),
        email_data.get("chair_buying_signals"),

        sales_potential.get("useful_for_product"),
        sales_potential.get("relevance"),
        sales_potential.get("reason"),
        sales_potential.get("sales_opportunity"),

        email_data.get("sender_email"),
        email_data.get("subject"),
        email_data.get("body"),
        email_data.get("personalization_point"),

        email_data.get("status"),
    ]


# ============================================================
# EXPORT
# ============================================================

def export_sales_results(
    emails: list[dict],
    filename: str | None = None,
) -> str:
    """
    Export final sales results to an Excel workbook.

    Returns the absolute path of the generated file.
    """

    if filename is None:

        timestamp = datetime.now().strftime(
            "%Y%m%d_%H%M%S"
        )

        filename = (
            f"sales_leads_{timestamp}.xlsx"
        )

    output_path = EXPORT_DIR / filename

    # --------------------------------------------------------
    # Workbook
    # --------------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Sales Leads"

    # --------------------------------------------------------
    # Headers
    # --------------------------------------------------------

    headers = [
        "Business Name",
        "Category",
        "City",
        "Phone",
        "Lead Email",
        "Website",

        "Qualification",
        "Qualification Reason",
        "Qualification Confidence",

        "Research Depth",
        "Business Summary",
        "Industry",
        "Business Activity",

        "Hiring Signals",
        "Expansion Signals",
        "Workplace Signals",
        "Office Chair Buying Signals",

        "Useful For Product",
        "Relevance",
        "Sales Reason",
        "Sales Opportunity",

        "Sender Email",
        "Email Subject",
        "Email Body",
        "Personalization",

        "Email Status",
    ]

    worksheet.append(headers)

    # --------------------------------------------------------
    # Header styling
    # --------------------------------------------------------

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    for email_data in emails:

        row = _build_row(
            email_data
        )

        row = [
            _to_text(value)
            for value in row
        ]

        worksheet.append(row)

    # --------------------------------------------------------
    # Formatting
    # --------------------------------------------------------

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # --------------------------------------------------------
    # Excel table
    # --------------------------------------------------------

    if worksheet.max_row >= 2:

        table_ref = (
            f"A1:"
            f"{get_column_letter(worksheet.max_column)}"
            f"{worksheet.max_row}"
        )

        table = Table(
            displayName="SalesLeadsTable",
            ref=table_ref,
        )

        style = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )

        table.tableStyleInfo = style

        worksheet.add_table(table)

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    widths = {
        "A": 28,
        "B": 22,
        "C": 18,
        "D": 18,
        "E": 30,
        "F": 35,

        "G": 16,
        "H": 40,
        "I": 24,

        "J": 16,
        "K": 45,
        "L": 20,
        "M": 22,

        "N": 40,
        "O": 40,
        "P": 40,
        "Q": 40,

        "R": 20,
        "S": 16,
        "T": 40,
        "U": 45,

        "V": 30,
        "W": 35,
        "X": 65,
        "Y": 45,

        "Z": 18,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------------
    # Wrap text
    # --------------------------------------------------------

    for row in worksheet.iter_rows():

        for cell in row:

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    # --------------------------------------------------------
    # Save
    # --------------------------------------------------------

    workbook.save(output_path)

    print(
        f"\nExcel exported successfully:"
        f"\n{output_path}"
    )

    return str(output_path)